"""casemap_service.py — the ONE CaseMap processing pipeline.

Extracted from server/app.py so there is a single source of truth for how a
document becomes graph nodes/edges. Both entry points are thin wrappers over
this module:

  * server/app.py  — HTTP wrapper (upload → JSON), adds no pipeline logic.
  * scripts/poc_run.py — CLI/batch wrapper (folder → Markdown proof report +
    poc_graph.json), adds only report formatting and the poc-only bi-temporal
    edge experiment.

Before this split the two had drifted badly: every 2026-09 improvement
(SaT boundary refinement, uncovered-date events, polarity reclassification,
act-name/provisions, the cross-document entity gate, section fallbacks,
rhetorical roles, word-amounts) lived ONLY in server/app.py, so poc_run.py
silently ran an inferior pipeline while claiming to prove "the pipeline."
This module ends that: change the pipeline here, both surfaces get it.

Nothing here is generative — every model used (OpenNyAI NER, SaT boundary
judge, GLiNER party judge) only SELECTS or CLASSIFIES existing text, never
writes new text (see THESIS §3 / FORBIDDEN C11). Model handles are injected
(ml_nlp) or lazily loaded as process-wide singletons; this module itself
never persists anything to disk — persistence is each wrapper's own choice.
"""

from __future__ import annotations

import os
import re
import traceback
from collections import Counter

import opennyai_bridge as bridge
import casemap_pipeline as pipeline
from document_profile import extract_parties_hybrid, party_result_to_fallback_event
from casemap_pipeline import (
    extract_pages, segment_document_layered, build_section_text,
    extract_deterministic, extract_entities, normalize_entities, detect_events,
    detect_events_srl,
    build_inverted_index, generate_candidate_pairs, score_pair,
    sparsify_and_cluster, label_edge, to_react_flow, make_similarity_fn,
    _classify_polarity, offset_to_page,
)
from important_lines import extract_important_lines
from case_symbols import provision_key, _ACT_IN_SPAN, find_word_amounts
from rhetorical_roles import tag_paragraphs, split_paragraphs

ACT_WINDOW_CHARS = 160  # how far around a bare "Section 22" hit to look for "... Act, 1956"
_TIGHT_ACT_RE = re.compile(r"([A-Z][\w&/,.\-]*(?:\s+[A-Z][\w&/,.\-]*){0,4}\s+Act(?:,\s*\d{4})?)$")

EMBED_MODEL = "all-MiniLM-L6-v2"
CONTEXT_CHARS = 450  # each side of a sentence, for the "expand to paragraph" view
EST_CHARS_PER_PAGE = 3000  # rough estimate for single-spaced 12pt legal text


# ---------------------------------------------------------------------------
# Model layer — lazy, process-wide singletons (loaded once, reused).
# ---------------------------------------------------------------------------
_ML_NLP = "_unloaded"


def get_ml_nlp():
    global _ML_NLP
    if _ML_NLP == "_unloaded":
        try:
            print("[+] loading OpenNyAI legal NER (mandatory layer) ...", flush=True)
            _ML_NLP = bridge.load_opennyai_ner("sm", allow_degraded=True)
            print("[+] ML layer:", "active" if _ML_NLP is not None else "DEGRADED", flush=True)
        except Exception:
            traceback.print_exc()
            _ML_NLP = None
    return _ML_NLP


_SAT_MODEL = "_unloaded"


def _get_sat_model():
    """segment-any-text/sat-3l-sm — a small (~0.2B) model fine-tuned
    specifically for sentence-boundary detection, robust on messy/corrupted
    text by design (exactly what OCR'd or Word-extracted legal text is).
    This is a JUDGE, not a generator: it only decides where to cut inside
    text that's already there — it never sees or produces the fact text
    itself, so it can't introduce a hallucinated fact, only a better (or,
    degraded, no worse than before) cut point. Loaded once, lazily, same
    pattern as get_ml_nlp(). Optional layer: falls back to the
    deterministic _line_bounds()/_sentence_bounds() if unavailable."""
    global _SAT_MODEL
    if _SAT_MODEL == "_unloaded":
        try:
            print("[+] loading sentence-boundary judge (sat-3l-sm) ...", flush=True)
            from wtpsplit_lite import SaT
            _SAT_MODEL = SaT("sat-3l-sm")
            print("[+] boundary judge: active", flush=True)
        except Exception:
            traceback.print_exc()
            print("[!] boundary judge unavailable — falling back to regex heuristics", flush=True)
            _SAT_MODEL = None
    return _SAT_MODEL


def _sat_sentence_spans(text: str) -> list[tuple[int, int]] | None:
    """Split `text` into real sentences via the SaT judge and return their
    (start, end) character offsets. None if the model isn't available."""
    sat = _get_sat_model()
    if sat is None:
        return None
    spans = []
    pos = 0
    for sent in sat.split(text):
        start, end = pos, pos + len(sent)
        pos = end
        lstripped = len(sent) - len(sent.lstrip())
        rstripped = len(sent) - len(sent.rstrip())
        spans.append((start + lstripped, end - rstripped))
    return spans


_GLINER_MODEL = "_unloaded"
_PARTY_LABELS = ["person name", "organization, company, or government body name"]
# TRIED 2026-09-11 (F-22 addendum) and REJECTED: adding a third "address,
# place name, or geographic location" label to catch address-continuation
# fragments that leak through as fake parties ("Khashewadi, Tiroda",
# "Grampanchayat Tiroda" on a real NGT appeal caption). It did catch one
# real case -- but zero-shot label sets are not independent: adding that
# third label changed this SAME model's span/label choice on text having
# nothing to do with addresses, silently shrinking "N. RAM" (a real
# petitioner, just fixed above) from a full 6/6-char match down to a 3/6
# partial ("RAM" only), enough to drop it below the keep threshold again.
# A second, single-label-only call to isolate the address signal was tried
# too and was worse: with no other label to compete against, GLiNER tagged
# "address" on almost everything tested, including "The Sarpanch" and "The
# District Collector" at full coverage. Both a real regression on a
# just-fixed case and a promiscuous single-label signal -- not shipped.
# The remaining address-fragment noise stays a known residual (fix it
# structurally in document_profile.py's entry continuation logic, not by
# asking this judge to type every fragment) rather than trade one real bug
# for another.


def _get_gliner_model():
    """urchade/gliner_small-v2.1 — a small zero-shot span classifier, used
    here purely as a JUDGE over what extract_parties_hybrid() already
    produced: keep or drop an already-extracted candidate name. It is
    never asked to propose new candidates and never rewrites the ones it's
    given — that keeps the verbatim guarantee intact (see FINDINGS.md F-11
    on why a GENERATIVE model was rejected for this pipeline; this is a
    different, narrower risk profile). Optional layer: on failure, party
    filtering falls back to the existing tier/confidence flagging alone."""
    global _GLINER_MODEL
    if _GLINER_MODEL == "_unloaded":
        try:
            print("[+] loading party-name judge (gliner_small-v2.1) ...", flush=True)
            from gliner import GLiNER
            _GLINER_MODEL = GLiNER.from_pretrained("urchade/gliner_small-v2.1")
            print("[+] party judge: active", flush=True)
        except Exception:
            traceback.print_exc()
            print("[!] party judge unavailable — falling back to tier/confidence flagging only", flush=True)
            _GLINER_MODEL = None
    return _GLINER_MODEL


# "& Ors."/"& Anr."/"and Others" is pure Indian-legal-caption boilerplate
# appended after a real name ("N. Ram & Ors") -- never itself part of what
# makes something a name. Found on a real filed writ petition (2026-09-11):
# GLiNER returns ZERO entities for "N. RAM & ORS" (score 0) but correctly
# tags "N. RAM" alone at 0.40 -- the suffix confuses the model into missing
# the real name attached to it. Stripped before judging only; the stored
# party name itself keeps the suffix (it's real information -- unnamed
# co-parties exist).
_ORS_SUFFIX_RE = re.compile(r"\s*(?:&|and)\s*(?:ors?|anrs?|others?)\.?\s*$", re.I)

# "State of <State>"/"Union of India"/"Government of <Place>" are the single
# most common respondent pattern in Indian litigation -- and, found while
# fixing the NGT appeal above (2026-09-11, F-22/F-23), GLiNER is erratic on
# it: "State of U.P." judges fine (0.41, full coverage) but "State of
# Maharashtra" and "State of Madhya Pradesh" -- textually the SAME
# pattern, just an unabbreviated state name -- return ZERO entities. Given
# how templated and unambiguous this specific pattern is, and that a small
# zero-shot model's inconsistency here would otherwise be re-discovered
# state-by-state, it gets a deterministic bypass instead of relying on the
# model to eventually get it right.
_GOVT_LITIGANT_RE = re.compile(
    r"^\s*(?:the\s+)?(?:state\s+of\s+[A-Z][A-Za-z.\s]{2,40}|"
    r"union\s+of\s+india|union\s+government|central\s+government|"
    r"state\s+government|government\s+of\s+[A-Z][A-Za-z.\s]{2,40})\s*$", re.I)


def _judge_party_name(name: str) -> bool:
    """True if the party judge thinks `name` is plausibly a real person/org
    name, False if it looks like extraction noise (a section heading, a
    stray phrase, a form field). If the judge is unavailable, default to
    keeping the name (fail open — no judge means no additional filtering,
    not silent data loss)."""
    judge_name = _ORS_SUFFIX_RE.sub("", name).strip() or name
    if _GOVT_LITIGANT_RE.match(judge_name):
        return True
    model = _get_gliner_model()
    if model is None:
        return True
    try:
        ents = model.predict_entities(judge_name, _PARTY_LABELS, threshold=0.3)
    except Exception:
        traceback.print_exc()
        return True
    # Keep when the judge's entities TOGETHER cover most of the candidate.
    # Not "any single entity" -- a real name the model splits into two
    # adjacent spans ("The Sarpanch" + "Grampanchayat Tiroda", a real NGT
    # appellant, 2026-09-11) would otherwise fail a per-entity 70% test
    # even though the two spans jointly cover 94% of the string. Overlaps
    # are naturally rare for span predictions over a single short candidate,
    # so a plain sum is a fine approximation of true coverage here.
    covered = sum(e["end"] - e["start"] for e in ents)
    return covered >= 0.7 * len(judge_name)


# Human-readable model identifiers, surfaced in layer_status() so a health
# check confirms not just "loaded" but *which* model/version is running --
# distinguishing e.g. a stale VPS deploy from the one just pushed is
# otherwise invisible from the outside.
MODEL_NAMES = {
    "ml_layer": "opennyai/en_legal_ner_sm 3.2.0",
    "boundary_judge": "segment-any-text/sat-3l-sm",
    "party_judge": "urchade/gliner_small-v2.1",
    "embedding_layer": EMBED_MODEL,
}


def layer_status() -> dict:
    """Health snapshot for the ML layers. Forces the ML NER to load (so its
    status is accurate) but reports the optional judges/embedder only if
    already touched -- unless warm_all_models() was run at startup, in which
    case every layer here is already loaded and this is a cheap read."""
    nlp = get_ml_nlp()
    embed_loaded = EMBED_MODEL in pipeline._EMBED_MODELS
    return {"status": "ok",
            "ml_layer": "active" if nlp is not None else "degraded",
            "boundary_judge": "active" if _SAT_MODEL not in ("_unloaded", None) else "unloaded",
            "party_judge": "active" if _GLINER_MODEL not in ("_unloaded", None) else "unloaded",
            "embedding_layer": "active" if embed_loaded else "unloaded",
            "models": MODEL_NAMES}


def warm_all_models() -> dict:
    """Force every lazy model layer to load right now, instead of each one
    paying its own load cost (and memory spike) on whichever request happens
    to touch it first. Meant for hardware with headroom to spare at boot --
    a VPS, not Render free tier's 512MB, which OOM-crashed mid-request when
    the judges + embedder loaded lazily under real traffic (2026-09-11).
    See server/app.py's WARM_MODELS_ON_START."""
    get_ml_nlp()
    _get_sat_model()
    _get_gliner_model()
    pipeline.warm_similarity_model(EMBED_MODEL)
    return layer_status()


# ---------------------------------------------------------------------------
# Ingestion — one record per page, text resolved per file type.
# ---------------------------------------------------------------------------
def _docx_to_pages(path: str, name: str) -> list[dict]:
    """Word has no fixed page grid the way a PDF does — pagination is
    computed at render time from fonts/margins, which python-docx can't
    see. Two signals combined, not one:

    1. Explicit page breaks the author inserted (Ctrl+Enter) — the
       strongest signal where present, but real documents don't break
       every page: checked on a real petition, only 9 explicit breaks for
       10 sections, and everything AFTER the last one (GROUNDS, PRAYER,
       AFFIDAVIT, ...) landed in one ~35,000-character block counted as a
       single "page" — obviously wrong; that's 10+ real printed pages.
    2. A character-count estimate (~3000 chars/page, single-spaced 12pt)
       to sub-divide any block too long to plausibly be one real page,
       splitting only at paragraph boundaries so nothing is cut mid-
       sentence. This is an ESTIMATE, not true rendering — python-docx
       cannot know the real page count without actually laying the
       document out (fonts, margins, tables all affect it) — but it is a
       real, substantial improvement over "one giant page," and every
       page number this produces is honestly an estimate, not a claim of
       exactness.
    """
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(path)

    pages_text: list[list[str]] = [[]]
    for p in d.paragraphs:
        has_break = any(
            br.get(qn("w:type")) == "page"
            for run in p.runs
            for br in run._element.findall(qn("w:br"))
        )
        pages_text[-1].append(p.text)
        if has_break:
            pages_text.append([])

    final_chunks: list[list[str]] = []
    for chunk in pages_text:
        current: list[str] = []
        current_len = 0
        for para in chunk:
            para_len = len(para) + 1
            if current and current_len + para_len > EST_CHARS_PER_PAGE:
                final_chunks.append(current)
                current, current_len = [], 0
            current.append(para)
            current_len += para_len
        final_chunks.append(current)

    pages = [{
        "document": name, "document_path": path, "page_number": i,
        "text": "\n".join(chunk), "page_kind": "digital", "words": [],
    } for i, chunk in enumerate(final_chunks, start=1) if "\n".join(chunk).strip()]

    return pages or [{
        "document": name, "document_path": path, "page_number": 1,
        "text": "\n".join(p.text for p in d.paragraphs), "page_kind": "digital", "words": [],
    }]


def _txt_to_pages(path: str, name: str) -> list[dict]:
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return [{
        "document": name, "document_path": path, "page_number": 1,
        "text": text, "page_kind": "digital", "words": [],
    }]


# ---------------------------------------------------------------------------
# Per-event enrichment helpers.
# ---------------------------------------------------------------------------
def _link_dates_in_place(evs: list[dict], dates: list[dict]) -> None:
    """IMPORTANT_LINE events (important_lines.py) and SECTION fallback
    events (below) never attempt date-linking at all — important_lines.py
    hardcodes linked_date=None regardless of whether the sentence it just
    extracted contains a real date ("BCI issued Circular No. 13/2024 ... on
    24.09.2024" is a real example: a genuine key point, with a genuine
    date, shown as "undated"). detect_events() only links a date to a
    KEYWORD-triggered event (ORDER/NOTICE/PAYMENT/...) within a fixed
    window of that keyword — a sentence with a date but no recognized
    keyword nearby gets no dated event either. Both are real gaps, not
    document limitations: `dates` here is extract_deterministic()'s own
    per-section date scan, already computed: if one of those dates sits
    inside (or right at the edge of, allowing for sentence-boundary
    slack) an event's own sentence span, attach it."""
    if not dates:
        return
    for e in evs:
        if e.get("linked_date"):
            continue
        s = e["sources"][0]
        cs, ce = s["char_start"], s["char_end"]
        candidates = [d for d in dates if cs - 20 <= d["span"][0] and d["span"][1] <= ce + 20]
        if candidates:
            e["linked_date"] = min(candidates, key=lambda d: abs(d["span"][0] - cs))


def _line_bounds(text: str, start: int, end: int) -> tuple[int, int] | None:
    """A "list of dates" section isn't prose — it's a real table, one row
    per Word paragraph, each row literally "DATE\tDescription" (confirmed
    directly: '21.07.2026\\tSupreme Court in W.P.(C) No. 31/2025 ... final
    disposal on 25.08.2026. Annexure P-8.' is ONE paragraph, ONE row,
    joined into section_text with real newlines by _docx_to_pages()).
    Guessing at sentence boundaries for this kind of content is the wrong
    tool — it can split one row into two fragments at whichever embedded
    period or date happens to fall mid-row (a real, reported example:
    21.07.2026 and 25.08.2026 both appear in the SAME row, and sentence-
    boundary guessing cut the row in half between them). The document's
    own real paragraph boundaries are a much better signal than a period
    search. Returns None (caller falls back to sentence bounds) when the
    line doesn't look like a real list row — too short to be meaningful,
    or long enough that it's more likely a normal wrapped paragraph than
    a table row."""
    lo = text.rfind("\n", 0, start)
    lo = lo + 1 if lo != -1 else 0
    hi = text.find("\n", end)
    hi = hi if hi != -1 else len(text)
    line = text[lo:hi].strip()
    if 20 <= len(line) <= 600:
        return lo, hi
    return None


def _sentence_bounds(text: str, start: int, end: int, max_len: int = 500) -> tuple[int, int]:
    """Expand (start, end) out to the nearest complete sentence on both
    sides — a reader should never see a snippet that begins or ends
    mid-word/mid-clause ("ver ification — not permission" was a real
    example). `max_len` bounds the search itself, not just the result: a
    real petition's cause-title clause runs 500+ characters with no period
    at all ("(A Writ Petition under Article 32 ... as: having no rational
    nexus ...)"), and searching further back for one just finds an
    unrelated EARLIER sentence, producing a snippet that still doesn't
    read as one clean unit. When no period exists inside the capped
    window, fall back to the nearest word boundary — worse than a real
    sentence edge, but never a mid-word cut, which is the one thing this
    must never do."""
    lo_floor = max(0, start - max_len)
    lo = text.rfind(". ", lo_floor, start)
    if lo != -1:
        lo = lo + 2
    else:
        sp = text.find(" ", lo_floor)
        lo = sp + 1 if 0 <= sp < start else lo_floor

    hi_ceiling = min(len(text), end + max_len)
    nxt = text.find(". ", end, hi_ceiling)
    if nxt != -1:
        hi = nxt + 1
    else:
        prev = text.rfind(" ", end, hi_ceiling)
        hi = prev if prev != -1 else hi_ceiling
    return lo, hi


def _refine_bounds_with_sat(events: list[dict], sat_spans: list[tuple[int, int]] | None,
                             section_text: str) -> None:
    """Layer SaT on top of whatever already chose these events' sentence
    spans (important_lines.py's own abbreviation-aware splitter, or
    detect_events()'s fixed keyword window) — WITHOUT changing what was
    selected. important_lines.py's splitter is a validated component (real
    bugs already found and fixed against the full test corpus — F-11/F-12,
    0 fragments across 52 real nodes) and stays the source of truth for
    WHICH sentence is important; this only tightens its edges when SaT's
    trained boundary disagrees (e.g. an abbreviation the regex splitter
    doesn't recognize). Ranking already happened on the ORIGINAL span
    before this runs, so refining the boundary afterward can't change
    which sentence was picked — only how cleanly it displays."""
    if not sat_spans:
        return
    for e in events:
        s = e["sources"][0]
        cs, ce = s["char_start"], s["char_end"]
        best, best_overlap = None, 0
        for ss, se in sat_spans:
            overlap = min(ce, se) - max(cs, ss)
            if overlap > best_overlap:
                best_overlap, best = overlap, (ss, se)
        if best is None or best == (cs, ce):
            continue
        ss, se = best
        # Guard against SaT merging several unrelated sentences into one
        # span (e.g. a short paragraph it fails to split at all) — only
        # accept a refinement that's in the same ballpark as the original.
        if se - ss > 2 * (ce - cs) + 200:
            continue
        new_text = section_text[ss:se].strip()
        if not new_text:
            continue
        s["char_start"], s["char_end"] = ss, se
        s["text"] = new_text
        if "section_text" in e:
            e["section_text"] = new_text
        display = re.sub(r"\s+", " ", new_text)
        e["label"] = display[:160] + ("…" if len(display) > 160 else "")


def _events_for_uncovered_dates(section_events: list[dict], dates: list[dict],
                                 sat_spans: list[tuple[int, int]] | None,
                                 section_text: str, section: dict, document_id: str) -> list[dict]:
    """_link_dates_in_place only attaches a date to a sentence that
    important_lines.py/detect_events() ALREADY turned into an event — but
    important_lines.py only keeps the top-few "most central" sentences per
    paragraph (embedding-centrality ranking, by design — see F-11). For an
    ordinary paragraph that's fine; for a genuine chronology section (LIST
    OF DATES AND EVENTS is the clear case — real example: 7 real dates in
    the section, only 0 of them survived into an event, because none of
    the sentences containing them happened to rank as "central") that
    silently drops most of the actual date list. Any real date not already
    covered by SOME event's sentence span gets its own node here, built
    directly from the text around the date itself — so the timeline
    reflects every date really in the document, not only the ones the
    relevance ranker happened to keep."""
    if not dates:
        return []
    covered_spans = [e["sources"][0] for e in section_events]

    def is_covered(d):
        return any(s["char_start"] <= d["span"][0] and d["span"][1] <= s["char_end"]
                   for s in covered_spans)

    # Computed once per section, reused for every date in it. Priority
    # order matters and was tested, not assumed: _line_bounds() (the
    # document's own real paragraph/row structure) goes FIRST when it
    # applies — a real regression was caught building this: the SaT judge,
    # tried first initially, is a general-purpose PROSE sentence segmenter,
    # and it misread the Indian legal abbreviation "SLP(Crl.)" as a
    # sentence end, cutting a real table row short in a way the structural
    # line-bounds approach didn't. _line_bounds only ever fires for a
    # short single-paragraph unit (<=600 chars — see its own docstring),
    # so it naturally defers to SaT for anything long enough to be real
    # wrapped prose, where SaT genuinely is the better judge (verified: it
    # correctly kept a table row with an embedded second date as one unit,
    # which the OLD sentence-only regex could not do at all). `sat_spans`
    # is computed once per section by the caller and passed in — shared
    # with _refine_bounds_with_sat() rather than recomputed here.
    out = []
    for d in dates:
        if is_covered(d):
            continue
        bounds = _line_bounds(section_text, d["span"][0], d["span"][1])
        if bounds is None and sat_spans:
            bounds = next(((s, e) for s, e in sat_spans
                          if s <= d["span"][0] and d["span"][1] <= e), None)
        if bounds is None:
            bounds = _sentence_bounds(section_text, d["span"][0], d["span"][1])
        lo, hi = bounds
        snippet = section_text[lo:hi].strip()
        display = re.sub(r"\s+", " ", snippet)
        out.append({
            "type": "DATED_EVENT",
            "document_id": document_id,
            "section_label": section["section_label"],
            "section_text": snippet,
            "label": display[:160] + ("…" if len(display) > 160 else ""),
            "linked_entities": [], "linked_date": d, "linked_amount": None,
            # Reuses detect_events()'s own polarity classifier on this
            # snippet's real text — previously hardcoded NEUTRAL regardless
            # of content, so a DATED_EVENT could never participate in
            # conflict detection even when its text plainly asserts or
            # denies something.
            "polarity": _classify_polarity(display.lower()), "confidence": "verbatim_fallback",
            "sources": [{
                "document": document_id, "page": d.get("page"),
                "char_start": lo, "char_end": hi,
                # No further length cut here — snippet is already bounded to
                # a clean sentence by _sentence_bounds(); re-slicing it with
                # a hard [:N] would reintroduce exactly the mid-word cut
                # this whole function exists to avoid. Display-level
                # truncation (previews in lists) is the frontend's job —
                # ui/app.js's previewText()/smartTruncate() handle that
                # against this full text, respecting sentence/word bounds.
                "text": snippet, "section": section["section_label"],
            }],
        })
    return out


def _attach_context(evs: list[dict], base_text: str) -> None:
    """Slice a ±CONTEXT_CHARS window around each event's primary source out
    of the SAME text its char_start/char_end were computed against. Passing
    the wrong base_text (e.g. the whole document instead of the section)
    silently produces a context window from an unrelated part of the
    document — see the caller in process_document for why this matters."""
    for e in evs:
        s = e["sources"][0]
        cs, ce = s["char_start"], s["char_end"]
        lo, hi = max(0, cs - CONTEXT_CHARS), min(len(base_text), ce + CONTEXT_CHARS)
        s["context"] = base_text[lo:hi]
        s["context_start"] = lo


def _trim_to_boundary(text: str, max_len: int) -> str:
    """Cut text to at most max_len chars, ending at a real sentence if one
    exists in range, else the last whole word — never mid-word. Used
    anywhere a highlighted/marked span in the UI comes from a hard
    character-count slice rather than the sentence-splitter."""
    if len(text) <= max_len:
        return text
    cut = text.rfind(". ", 0, max_len)
    if cut > max_len * 0.4:
        return text[:cut + 1]
    sp = text.rfind(" ", 0, max_len)
    return text[:sp] if sp > max_len * 0.4 else text[:max_len]


def _attach_rhetorical_roles(events: list[dict], section_text: str) -> None:
    """Label each event with the rhetorical role (FACTS / ISSUES / ARGUMENTS /
    ANALYSIS / RULING / PRECEDENT) of the paragraph it sits in, via the
    deterministic cue tagger (rhetorical_roles.py). Like the SaT and party
    judges, this only CLASSIFIES existing paragraphs — it never rewrites or
    invents text, so it can't touch the verbatim guarantee. Written onto
    sources[0] (which to_react_flow passes through unchanged) so the evidence
    drawer can show WHERE in a document's argument a fact comes from.
    Confidence is carried too: 'cue_matched' means a real role cue is in that
    paragraph; 'carried_forward' means it was inherited from an earlier
    tagged paragraph (weaker, and labelled as such in the UI)."""
    tagged = tag_paragraphs(split_paragraphs(section_text))
    if not tagged:
        return
    for e in events:
        cs = e["sources"][0]["char_start"]
        for t in tagged:
            if t["role"] and t["char_start"] <= cs < t["char_end"]:
                e["sources"][0]["rhetorical_role"] = t["role"]
                e["sources"][0]["rhetorical_confidence"] = t["confidence"]
                break


def _section_fallback_event(section: dict, section_text: str, document_id: str) -> dict:
    """One verbatim node for a detected heading that produced zero real
    events (short PRAYER/INTERIM RELIEF/QUESTIONS OF LAW clauses are common
    victims — too few sentences for important_lines, no EVENT_KEYWORDS hit
    for detect_events). Without this the heading is detected correctly by
    the backend but never appears anywhere in the UI."""
    label = section["section_label"]
    text = section_text.strip()
    display = re.sub(r"\s+", " ", text)
    page = section["pages"][0]["page_number"] if section.get("pages") else None
    return {
        "type": "SECTION",
        "document_id": document_id,
        "section_label": label,
        "section_text": text,
        "label": f"[{label}] " + (display[:160] + "…" if len(display) > 160 else display),
        "linked_entities": [],
        "linked_date": None,
        "linked_amount": None,
        # Same reasoning as _events_for_uncovered_dates() above — classify
        # against this section's own real text instead of a hardcoded
        # NEUTRAL, so a short PRAYER/AFFIDAVIT clause that plainly asserts
        # or denies something can actually take part in conflict detection.
        "polarity": _classify_polarity(display.lower()),
        "confidence": "verbatim_fallback",
        "sources": [{
            "document": document_id, "page": page,
            "char_start": 0, "char_end": len(text),
            "text": _trim_to_boundary(text, 400), "section": label,
        }],
    }


def _find_act_name(window: str) -> str | None:
    """Look for a real Act name in a prose window around a bare "Section N"
    citation. _ACT_IN_SPAN (case_symbols.py) is compiled case-INsensitive —
    correct for the tight, already-clean citation strings it was written
    for ("Section 7 of the Advocates Act, 1961"), but scanning open prose
    with it surfaces false positives: "no fact", "disciplinary act", "in
    charact[er]" all match "act" as a case-insensitive substring. Rather
    than change the shared regex (other callers rely on its current
    behavior on short strings, where this doesn't come up), require the
    literal capitalized word "Act" in what it found before trusting it —
    real Act names are always capitalized in these documents; ordinary
    prose using the word "act" is not."""
    m = _ACT_IN_SPAN.search(window)
    if not m:
        return None
    candidate = re.sub(r"\s+", " ", m.group("act")).strip()
    if len(candidate) > 120 or not re.search(r"\bAct\b", candidate):
        return None
    # _ACT_IN_SPAN's own non-greedy match can still walk back through an
    # entire lowercase clause to whatever capital letter starts the window
    # ("within the statutory scheme of the Advocates Act") — trim to just
    # the run of capitalized words immediately before "Act" so the UI shows
    # "Advocates Act" rather than the whole clause.
    tight = _TIGHT_ACT_RE.search(candidate)
    return tight.group(1) if tight else candidate


def _normalize_act_for_key(act: str | None) -> str | None:
    """"UGC Act" and "UGC Act, 1956" are the same Act, but provision_key()
    (case_symbols.py) keys on the act-name text verbatim, so a document
    that names the year sometimes and not other times (real, common —
    Indian filings often drop the year on later mentions) split into
    separate "Section 22" entries in the counsel report: one line for
    each act-name spelling. Strip a trailing ", YYYY" before keying only —
    the fuller, year-bearing spelling is still what gets displayed."""
    if not act:
        return None
    return re.sub(r",?\s*\d{4}\s*$", "", act).strip()


def _build_provisions(doc_results: list[dict]) -> list[dict]:
    """Cross-document provision table: "Section 22" cited in two documents
    should read as one entry, WITH the Act it belongs to when a document's
    own text names it nearby — case_symbols.provision_key() is the existing
    key function for this (act name + section number, or section-only when
    no Act is stated close enough to trust). No act-name inheritance across
    sentences that don't name it (see case_symbols.py's own docstring on
    why that's deliberately out of scope) — a bare "Section 7" stays
    grouped only with other bare "Section 7" mentions, never guessed onto
    whichever Act happened to be named earlier in the document."""
    groups: dict[str, dict] = {}
    for doc in doc_results:
        for prov in doc.get("provisions_detail") or []:
            norm_act = _normalize_act_for_key(prov["act"])
            text_for_key = f"{prov['raw']} of the {norm_act}" if norm_act else prov["raw"]
            key = provision_key(text_for_key)
            if not key:
                continue
            g = groups.setdefault(key, {"section": prov["raw"], "act": prov["act"],
                                         "documents": set(), "count": 0,
                                         "statute_lookup": None})
            g["count"] += 1
            g["documents"].add(doc["name"])
            # Prefer the fullest spelling seen (the one WITH a year, when
            # one exists) over strictly first-seen.
            if prov["act"] and (not g["act"] or len(prov["act"]) > len(g["act"])):
                g["act"] = prov["act"]
            if prov.get("statute_lookup") and not g["statute_lookup"]:
                g["statute_lookup"] = prov["statute_lookup"]
    return [{
        "key": key, "section": g["section"], "act": g["act"],
        "documents": sorted(g["documents"]), "count": g["count"],
        "statute_lookup": g["statute_lookup"],
    } for key, g in groups.items()]


# ---------------------------------------------------------------------------
# The pipeline — one document in, one result dict out.
# ---------------------------------------------------------------------------
def process_document(path: str, name: str, ml_nlp, ocr_engine: str = "tesseract",
                     ocr_lang: str = "eng", force_ocr: bool = False,
                     lookup_provisions: bool = False) -> dict:
    """Full per-document extraction. Returns a result dict whose "events" key
    holds the graph nodes for this document; the rest is the per-document
    summary the UI and the report both read. `ml_nlp` is injected (loaded
    once by the caller); the SaT/GLiNER judges lazy-load as module singletons.

    `lookup_provisions` (default False, opt-in): fetch each cited
    provision's own text from IndianKanoon.org (provision_lookup.py) — one
    outbound call per UNIQUE provision in this document, query is the bare
    citation string only, never document content. Off unless the caller
    explicitly asks, since it's the one part of this pipeline that reaches
    the internet at all.
    """
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf":
        result = extract_pages(path, ocr_engine=ocr_engine, ocr_lang=ocr_lang,
                               force_ocr=force_ocr)
        pages, toc = result["pages"], result["toc"]
        sections, struct_meta = segment_document_layered(path, pages, toc)
    elif ext in (".docx", ".doc"):
        pages = _docx_to_pages(path, name)
        sections, struct_meta = segment_document_layered(path, pages, [])
    else:  # .txt
        pages = _txt_to_pages(path, name)
        sections, struct_meta = segment_document_layered(path, pages, [])

    page_kinds = Counter(p["page_kind"] for p in pages)
    full_text = "\n".join(p["text"] for p in pages)

    party_result = extract_parties_hybrid(full_text, name, nlp=ml_nlp, allow_degraded=True)

    events: list[dict] = []
    doc_paragraph_no = 1
    all_sections_cited: list[str] = []
    provisions_detail: list[dict] = []
    all_amounts: list[str] = []
    all_case_numbers: list[str] = []

    for section in sections:
        section_text, offsets = build_section_text(section)
        if not section_text.strip():
            continue
        det = extract_deterministic(section_text, offsets)
        # Amounts written in words ("Rs. Fifty Lakhs Only", "rupees one
        # crore") are invisible to the digit-only figure regex — add them so
        # coverage matches what a reader sees. Skipped when a figure amount
        # already sits at the same spot, so a "Rs. 50,00,000 (Rupees Fifty
        # Lakhs)" pair isn't double-counted.
        _fig_amt_starts = [a["span"][0] for a in det["amounts"]]
        for _wa in find_word_amounts(section_text):
            if any(abs(_wa["span"][0] - s) < 40 for s in _fig_amt_starts):
                continue
            det["amounts"].append({
                "raw": _wa["raw"], "span": _wa["span"],
                "page": offset_to_page(_wa["span"][0], offsets)})
        ents = extract_entities(section_text, offsets)
        ent_map = normalize_entities(ents)
        evs = detect_events(section, section_text, offsets, det, ent_map, ents, name)

        il_evs, doc_paragraph_no = extract_important_lines(
            section, section_text, offsets, name,
            start_paragraph_no=doc_paragraph_no, use_embeddings=True)

        section_events = evs + il_evs
        # Computed once per section, shared by the boundary refinement
        # below and the uncovered-dates fallback — one real sentence split
        # of this section's text from the trained boundary judge.
        sat_spans = _sat_sentence_spans(section_text)
        _refine_bounds_with_sat(section_events, sat_spans, section_text)

        # Polarity (ASSERTS/DENIES/NEUTRAL — what conflict detection keys
        # off) is reclassified HERE, uniformly, for every event, using each
        # event's own (now SaT-tightened) sentence text — not
        # detect_events()'s original ±400-char display window. Found by
        # direct testing: that wide window is fine for readability but far
        # too wide for polarity — in anything but a long document it spans
        # MULTIPLE unrelated sentences, so a "denies" in one sentence
        # silently classified every OTHER nearby event (AGREEMENT, FILING,
        # ...) as DENIES too, a real false-conflict risk. important_lines.py
        # separately hardcodes NEUTRAL regardless of content, which this
        # also replaces, so a Key Fact that plainly asserts or denies
        # something can now actually take part in conflict detection.
        for e in section_events:
            e["polarity"] = _classify_polarity(e["sources"][0]["text"].lower())

        _link_dates_in_place(section_events, det.get("dates") or [])
        section_events += _events_for_uncovered_dates(
            section_events, det.get("dates") or [], sat_spans, section_text, section, name)

        # SRL layer (F-16): WHO/ACTION/WHAT/WHEN read off the dependency
        # parse for any date- or modal-bearing sentence none of the layers
        # above already turned into an event. covered_spans is everything
        # found so far for this section -- this only adds what's still
        # missing (measured at 23% of real event sentences on real
        # judgments, see scripts/poc_srl_events_v2.py).
        covered_spans = [(e["sources"][0]["char_start"], e["sources"][0]["char_end"])
                         for e in section_events]
        srl_evs = detect_events_srl(section, section_text, offsets, covered_spans,
                                     ent_map, ents, name)
        _link_dates_in_place(srl_evs, det.get("dates") or [])
        section_events += srl_evs

        # Every event above carries char_start/char_end relative to THIS
        # section's own text (see important_lines.py:272-273,
        # detect_events()'s `idx` in casemap_pipeline.py) — not the whole
        # document. Slicing the whole-document full_text with those numbers
        # (the original bug here) grabs text from a completely unrelated
        # part of the document once a section starts more than a few
        # hundred characters into the file. Context must come from the same
        # section_text the offsets were computed against.
        _attach_context(section_events, section_text)

        if not section_events:
            # A real, detected heading (PRAYER, INTERIM RELIEF, QUESTIONS OF
            # LAW, ...) that happens to be too short/keyword-free to produce
            # a KEY_FACT or important-line event would otherwise vanish from
            # the UI entirely, even though the backend found it. Surface it
            # verbatim instead of dropping it.
            section_events = [_section_fallback_event(section, section_text, name)]
            _link_dates_in_place(section_events, det.get("dates") or [])
            _attach_context(section_events, section_text)

        # Rhetorical role runs last, after the event set (normal OR fallback)
        # is final and its char offsets are settled — every event here carries
        # a section_text-relative char_start the tagger can locate.
        _attach_rhetorical_roles(section_events, section_text)
        events.extend(section_events)

        for hit in det.get("sections") or []:
            all_sections_cited.append(hit["raw"])
            # "Section 22" alone never names an Act — look at the text
            # around the citation for "... Act, 1956" the way a reader
            # would. _ACT_IN_SPAN is the same regex case_symbols.py uses to
            # key cross-document provisions; reused here (not reimplemented)
            # so "Section 22 of the UGC Act" and a later bare "Section 22"
            # in the same Act's context key the same way.
            start, end = hit["span"]
            lo, hi = max(0, start - ACT_WINDOW_CHARS), min(len(section_text), end + ACT_WINDOW_CHARS)
            window = section_text[lo:hi]
            act = _find_act_name(window)
            provisions_detail.append({
                "raw": hit["raw"], "act": act,
                "page": hit.get("page"), "window": window.strip(),
            })

        all_amounts.extend(r["raw"] for r in det.get("amounts") or [])
        all_case_numbers.extend(r["raw"] for r in det.get("case_numbers") or [])

    if lookup_provisions and provisions_detail:
        from provision_lookup import lookup_provision
        cache: dict[tuple[str, str], dict | None] = {}
        for p in provisions_detail:
            key = (p["raw"].strip().lower(), (p["act"] or "").strip().lower())
            if key not in cache:
                cache[key] = lookup_provision(p["raw"], p["act"])
            p["statute_lookup"] = cache[key]

    fb = party_result_to_fallback_event(party_result, name, [(0, len(full_text), 1)])
    if fb is not None:
        _attach_context([fb], fb["section_text"])
        events.append(fb)

    return {
        "name": name,
        "pages": len(pages),
        "page_kinds": dict(page_kinds),
        "structure_tier": struct_meta["tier"],
        "structure_confidence": struct_meta["confidence"],
        "structure_reason": struct_meta["tier_reason"],
        "sections": len(sections),
        "events": events,
        # Party judge (GLiNER) runs after extract_parties_hybrid(), not
        # instead of it — it only decides keep/drop on names the existing
        # ML+deterministic layers already produced, never proposes a name
        # of its own. A dropped name never silently vanishes without
        # trace: party_tier/party_confidence below still reflect the
        # underlying extraction's own confidence regardless of what the
        # judge did.
        "parties": [{"name": p.name, "role": p.role, "confidence": p.role_confidence}
                    for p in party_result.parties if _judge_party_name(p.name)],
        "party_tier": party_result.tier,
        "party_confidence": party_result.confidence,
        "sections_cited": sorted(set(all_sections_cited)),
        "provisions_detail": provisions_detail,
        "amounts": sorted(set(all_amounts)),
        "case_numbers": sorted(set(all_case_numbers)),
        "char_count": len(full_text),
    }


# ---------------------------------------------------------------------------
# The graph — many document results in, one react-flow graph out.
# ---------------------------------------------------------------------------
def build_case_graph(doc_results: list[dict], all_events: list[dict],
                     embed_model: str = EMBED_MODEL, *, return_details: bool = False):
    """Assemble the cross-document event graph from per-document events.

    `doc_results` are the process_document() summaries WITH "events" already
    popped into `all_events` (matching how both wrappers call this). Returns
    the react-flow graph dict {nodes, edges, documents, provisions}; the
    caller adds any runtime/ml-layer fields it wants. With return_details=True
    also returns the scoring internals (kept edges, score breakdowns, edge
    labels, candidate count) so a caller can build extra reporting (the CLI's
    cross-doc/bi-temporal sections) without re-scoring.
    """
    provisions = _build_provisions(doc_results)
    if not all_events:
        graph = {"nodes": [], "edges": [], "documents": doc_results,
                 "provisions": provisions}
        if return_details:
            return graph, {"kept": set(), "scores": {}, "edge_labels": {},
                           "candidate_count": 0}
        return graph

    index = build_inverted_index(all_events)
    candidates = generate_candidate_pairs(index)
    sim_fn = make_similarity_fn(embed_model)
    n = len(all_events)
    scores = {p: score_pair(all_events[p[0]], all_events[p[1]], index, n, sim_fn)
              for p in candidates}
    kept, _clusters = sparsify_and_cluster(all_events, candidates, scores)

    # Cross-document edges require a real shared entity (party, org, or
    # other named identity — case_symbols-normalized, stable since the
    # normalize_entities() ID fix) as a hard gate, not just one weighted
    # term among several a high semantic_score alone could satisfy. Verified
    # directly: without this gate, generic legal boilerplate similarity
    # connected a civil property case to an unrelated financial cheque-bounce
    # case with zero shared parties. Same-document edges are unaffected —
    # they're already known related by being the same filing.
    kept = {(a, b) for (a, b) in kept
            if all_events[a]["document_id"] == all_events[b]["document_id"]
            or set(all_events[a]["linked_entities"]) & set(all_events[b]["linked_entities"])}

    edge_labels = {}
    for (a, b) in sorted(kept):
        edge_labels[(a, b)] = label_edge(all_events[a], all_events[b])

    graph = to_react_flow(all_events, kept, edge_labels)
    graph["documents"] = doc_results
    graph["provisions"] = provisions
    if return_details:
        return graph, {"kept": kept, "scores": scores, "edge_labels": edge_labels,
                       "candidate_count": len(candidates)}
    return graph
